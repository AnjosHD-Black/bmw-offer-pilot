import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


CODE_RE = re.compile(r"(?<![A-Z0-9])([A-Z0-9]{3,4})(?![A-Z0-9])")
COMMON_NOISE = {"BMW", "TEXT", "PAKET", "XYZ", "AUTO", "CAR"}


def _normalize_code(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None

    code = str(value).strip().upper()
    if re.fullmatch(r"[A-Z0-9]{3,4}", code):
        return code

    return None


def _extract_code_from_material(value: Optional[str], known_codes: Optional[List[str]] = None) -> Optional[str]:
    if value is None:
        return None

    text = str(value).strip().upper()
    if not text:
        return None

    matches = CODE_RE.findall(text)
    if not matches:
        return None

    known = {str(code).upper() for code in (known_codes or [])}
    if known:
        for candidate in matches:
            if candidate in known:
                return candidate

    for candidate in matches:
        if candidate in COMMON_NOISE:
            continue
        if re.fullmatch(r"[A-Z0-9]{3,4}", candidate):
            return candidate

    return None


def _clean_price(value) -> Optional[str]:
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return str(value).replace('.', ',')

    text = str(value).strip()
    if not text:
        return None

    return text.replace('.', ',')


def parse_excel_option_data(file_path: str | Path, known_codes: Optional[List[str]] = None) -> Dict[str, object]:
    """
    Erwartetes Excel-Template:
    - Sheet: Data
    - Zeile 1 = Überschriften
    - Spalten: Material, Positionsbezeichung, Nettopreis, Währung

    Beispiel:
    Material: 'BMW 3AB Zusatzpaket'
    Positionsbezeichung: 'Sitzheizung'
    Nettopreis: 500
    Währung: EUR

    Wenn known_codes gesetzt ist, wird der Code bevorzugt aus dieser Menge gewählt,
    damit generische Wörter wie BMW oder TEXT nicht fälschlich als Code erkannt werden.
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb["Data"]

    headers = {}
    for idx, cell in enumerate(ws[1], start=1):
        headers[str(cell.value).strip()] = idx

    required_headers = ["Material", "Positionsbezeichung", "Nettopreis", "Währung"]
    missing = [name for name in required_headers if name not in headers]
    if missing:
        raise ValueError(f"Fehlende Excel-Spalten: {missing}")

    material_col = headers["Material"]
    label_col = headers["Positionsbezeichung"]
    price_col = headers["Nettopreis"]
    currency_col = headers["Währung"]

    all_codes: List[str] = []
    seen_codes = set()
    priced_lines: List[str] = []
    currency = "EUR"

    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < max(material_col, label_col, price_col, currency_col):
            continue

        material_value = row[material_col - 1]
        label_value = row[label_col - 1]
        price_value = row[price_col - 1]
        currency_value = row[currency_col - 1]

        code = _extract_code_from_material(material_value, known_codes=known_codes)
        if not code:
            continue

        code = code.upper()
        if code not in seen_codes:
            seen_codes.add(code)
            all_codes.append(code)

        label = str(label_value).strip() if label_value not in (None, "") else code
        price_text = _clean_price(price_value)
        if currency_value not in (None, ""):
            currency = str(currency_value).strip().upper()

        if price_text is None:
            continue

        priced_lines.append(f"{code} {label} {price_text}")

    return {
        "all_codes": all_codes,
        "priced_lines": priced_lines,
        "currency": currency,
    }
