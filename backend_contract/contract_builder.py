import os
from datetime import datetime

from docx import Document


def _add_key_value(document: Document, label: str, value: str):
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value or "-")


def build_contract(contract_data: dict) -> str:
    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(base_dir, "output")
    os.makedirs(output_dir, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"vertrag_{timestamp}.docx"
    file_path = os.path.join(output_dir, filename)

    document = Document()
    special_equipment = contract_data.get("special_equipment", [])
    buyer_name = " ".join(
        part for part in [
            contract_data.get("buyer_first_name", ""),
            contract_data.get("buyer_surname", ""),
        ]
        if part
    )

    document.add_heading("Kaufvertrag", level=1)
    document.add_paragraph(f"Erstellt am: {datetime.now().strftime('%d.%m.%Y')}")

    document.add_heading("Verkäufer", level=2)
    _add_key_value(document, "Verkäufer", contract_data.get("seller_code", ""))

    document.add_heading("Buyer", level=2)
    _add_key_value(document, "Name", buyer_name)
    _add_key_value(document, "Straße / Hausnummer", contract_data.get("buyer_street", ""))
    _add_key_value(document, "Stadt", contract_data.get("buyer_city", ""))
    _add_key_value(document, "Postleitzahl", contract_data.get("buyer_postal_code", ""))
    _add_key_value(document, "Country", contract_data.get("buyer_country", ""))
    _add_key_value(document, "VAT", contract_data.get("buyer_vat", ""))
    _add_key_value(document, "Tax No.", contract_data.get("buyer_tax_no", ""))
    _add_key_value(document, "Phone", contract_data.get("buyer_phone", ""))
    _add_key_value(document, "Mail", contract_data.get("buyer_email", ""))
    _add_key_value(document, "Represented by", contract_data.get("buyer_represented_by", ""))

    document.add_heading("Fahrzeugdaten", level=2)
    _add_key_value(document, "Baureihe", contract_data.get("contract_series", ""))
    _add_key_value(document, "Quantity", contract_data.get("quantity", ""))
    _add_key_value(document, "Price", contract_data.get("price", ""))
    _add_key_value(document, "FIN / VIN", contract_data.get("vin", ""))
    _add_key_value(document, "Order Nummer", contract_data.get("order_number", ""))
    _add_key_value(document, "Special Equipment", ", ".join(special_equipment) if special_equipment else "-")

    document.add_paragraph()
    document.add_paragraph("______________________________")
    document.add_paragraph("Unterschrift Kunde")
    document.add_paragraph()
    document.add_paragraph("______________________________")
    document.add_paragraph("Unterschrift Verkäufer")

    document.save(file_path)
    return file_path
