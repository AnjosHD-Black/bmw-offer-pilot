import os
from datetime import datetime

from docx import Document


def build_contract_g73(contract_data: dict) -> str:
    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(base_dir, "output")
    os.makedirs(output_dir, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"vertrag_g73_{timestamp}.docx"
    file_path = os.path.join(output_dir, filename)

    document = Document()
    document.add_heading("Kaufvertrag G73", level=1)
    document.add_paragraph(f"Erstellt am: {datetime.now().strftime('%d.%m.%Y')}")
    document.add_paragraph("G73-Vertragslayout wird separat gepflegt.")

    document.add_paragraph()
    document.add_paragraph("Order Nummer: " + (contract_data.get("order_number", "-") or "-"))
    document.add_paragraph("VIN: " + (contract_data.get("vin", "-") or "-"))

    document.save(file_path)
    return file_path
