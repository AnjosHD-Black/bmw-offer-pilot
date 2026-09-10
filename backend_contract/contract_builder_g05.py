import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


# ──────────────────────────────────────────────
# HELPERS
# ──────────────────────────────────────────────

def _plain(document: Document, text: str, size_pt: int = 11, bold: bool = False, space_before: int = 0, space_after: int = 0):
    """Single left-aligned paragraph, no heading style."""
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text or "")
    run.font.size = Pt(size_pt)
    run.bold = bold
    return p


def _spacer(document: Document, lines: int = 1):
    for _ in range(lines):
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def _add_header_image_slot(cell, image_path: str, placeholder: str, width_cm: float, alignment):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    if os.path.exists(image_path):
        run = paragraph.add_run()
        run.add_picture(image_path, width=Cm(width_cm))
        return

    run = paragraph.add_run(placeholder)
    run.bold = True


def _remove_table_borders(table):
    """Make a table completely invisible (no borders)."""
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = tcPr.find(qn("w:tcBorders"))
            if tcBorders is None:
                from docx.oxml import OxmlElement
                tcBorders = OxmlElement("w:tcBorders")
                tcPr.append(tcBorders)
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                from docx.oxml import OxmlElement
                border = OxmlElement(f"w:{side}")
                border.set(qn("w:val"), "none")
                tcBorders.append(border)


# ──────────────────────────────────────────────
# MAIN BUILD FUNCTION
# ──────────────────────────────────────────────

def build_contract_g05(contract_data: dict) -> str:
    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(base_dir, "output")
    assets_dir = os.path.join(base_dir, "assets")
    os.makedirs(output_dir, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"vertrag_g05_{timestamp}.docx"
    file_path = os.path.join(output_dir, filename)

    document = Document()
    # Slightly narrower margins for more usable width
    section = document.sections[0]
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(2.0)

    buyer_name = " ".join(
        part for part in [
            contract_data.get("buyer_first_name", ""),
            contract_data.get("buyer_surname", ""),
        ] if part
    )
    plz_city = " ".join(
        part for part in [
            contract_data.get("buyer_postal_code", ""),
            contract_data.get("buyer_city", ""),
        ] if part
    )
    # ── LOGO ROW ──────────────────────────────
    logo_table = document.add_table(rows=1, cols=2)
    logo_table.autofit = True
    _remove_table_borders(logo_table)
    left_cell, right_cell = logo_table.rows[0].cells

    _add_header_image_slot(
        left_cell,
        os.path.join(assets_dir, "bmw_group_wordmark.png"),
        "[BMW GROUP]",
        3.0,
        WD_ALIGN_PARAGRAPH.LEFT,
    )
    _add_header_image_slot(
        right_cell,
        os.path.join(assets_dir, "bmw_logo.png"),
        "[BMW LOGO]",
        2.0,
        WD_ALIGN_PARAGRAPH.RIGHT,
    )

    _spacer(document, 1)

    # ── RETURN ADDRESS LINE ───────────────────
    _plain(document, "BMW AG, 80000 München", size_pt=9, bold=False)

    _spacer(document, 1)

    # ── BUYER ADDRESS BLOCK (no labels) ───────
    _plain(document, buyer_name, size_pt=11, bold=False)
    _plain(document, contract_data.get("buyer_street", ""), size_pt=11)
    _plain(document, plz_city, size_pt=11)
    _plain(document, contract_data.get("buyer_country", ""), size_pt=11)
    _plain(document, f"Email: {contract_data.get('buyer_email', '')}", size_pt=10)
    _plain(document, f"Telephone Number: {contract_data.get('buyer_phone', '')}", size_pt=10)

    _spacer(document, 2)

    # ── REPRESENTED BY ────────────────────────
    rep = contract_data.get("buyer_represented_by", "")
    if rep:
        _plain(document, f"Represented by {rep}", size_pt=11)

    document.save(file_path)
    return file_path
